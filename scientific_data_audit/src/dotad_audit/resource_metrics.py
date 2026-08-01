from __future__ import annotations

import csv
import io
import re
import zipfile
from collections.abc import Iterable, Iterator, Sequence
from itertools import chain
from pathlib import Path
from typing import Any, TextIO

from openpyxl import load_workbook


PLACEHOLDER_TOKENS = {"", "na", "n/a", "nan", "null", "none", "nd", "--"}
ENTRY_ID_HEADERS = {"antibodyid", "entryid", "id"}
AMINO_ACID_ALPHABET = set("ACDEFGHIKLMNPQRSTVWYBXZJUO")
WORKBOOK_SEQUENCE_HEADERS = {"vh", "vl", "hcproteinsequence", "lcproteinsequence"}
HEAVY_SEQUENCE_HEADERS = (
    "abheavychainseq",
    "abheavyseq",
    "heavy",
    "hc",
    "vh",
    "vhhsequence",
    "vhsequence",
    "hcproteinsequence",
)
LIGHT_SEQUENCE_HEADERS = (
    "ablightchainseq",
    "ablightseq",
    "light",
    "lc",
    "vl",
    "vlsequence",
    "lcproteinsequence",
)
SINGLE_SEQUENCE_HEADERS = ("sequence",)
CDR_SEQUENCE_HEADER_GROUPS = (
    ("hcdr1", "cdrh1"),
    ("hcdr2", "cdrh2"),
    ("hcdr3", "cdrh3"),
    ("lcdr1", "cdrl1"),
    ("lcdr2", "cdrl2"),
    ("lcdr3", "cdrl3"),
)
BINDING_SEQUENCE_HEADERS = {
    "sequence",
    "vhhsequence",
    "vhsequence",
    "vlsequence",
    "abheavychainseq",
    "ablightchainseq",
    "abheavyseq",
    "ablightseq",
    "heavy",
    "light",
    "vh",
    "vl",
    "hc",
    "lc",
    "hcdr1",
    "hcdr2",
    "hcdr3",
    "lcdr1",
    "lcdr2",
    "lcdr3",
    "cdrh1",
    "cdrh2",
    "cdrh3",
    "cdrl1",
    "cdrl2",
    "cdrl3",
}
BINDING_META_HEADERS = {
    "antibodyid",
    "antibodyname",
    "abname",
    "name",
    "candidate",
    "poi",
    "target",
    "antigen",
    "agname",
    "agnamedetails",
    "agepitoperestrictions",
    "agpdbid",
    "agseq",
    "referenceab",
    "method",
    "format",
    "source",
    "sourcesampleid",
    "hcdrsdesigned",
    "design",
    "hcmut",
    "lcmut",
    "length",
    "editdistancetotrastuzumab",
    "abpdbid",
    "abstructuremethod",
    "agstructuremethod",
    "boundabagstructuremethod",
}
BINDING_SUPPORT_PATTERNS = (
    " std",
    "sd",
    "stddev",
    "stdev",
    "replicate",
    "replicates",
    "count",
    "counts",
)
BINDING_METRIC_KEYWORDS = (
    "kd",
    "affinity",
    "ic50",
    "ec50",
    "spr",
    "fitness",
    "expression",
    "enrichment",
    "tm",
    "stability",
    "psr",
    "sec",
    "smac",
    "hic",
    "hac",
    "acsins",
    "sins",
    "cic",
    "purity",
    "impurities",
    "pdi",
    "hmw",
    "lmw",
    "main",
    "score",
    "pred_affinity",
    "binding",
    "adcc",
    "recycling",
    "retention",
    "rt",
    "response",
    "heparin",
    "polyspecificity",
    "developability",
    "quality observation",
)
EXPERIMENTAL_META_HEADERS = {
    "antibodyid",
    "antibodyname",
    "hcsubtype",
    "lcsubtype",
    "productionbatch",
    "run",
    "technicalreplicate",
}
EXPERIMENTAL_SUPPORT_PATTERNS = (
    "stddev",
    "stdev",
    "replicate",
    "replicates",
    "count",
    "counts",
)
LITERATURE_META_HEADERS = {
    "antibodyid",
    "antibodyname",
    "antibody",
    "candidate",
    "vh",
    "vl",
    "heavy",
    "light",
    "chain",
    "mutation",
    "format",
    "units",
    "assayprotocol",
    "origin",
    "sourcesampleid",
    "hcdr1",
    "hcdr2",
    "hcdr3",
    "hcdrsdesigned",
    "cdrh1",
    "cdrh2",
    "cdrh3",
    "cdrl1",
    "cdrl2",
    "cdrl3",
    "target",
    "referenceab",
    "method",
    "design",
    "hcmut",
    "lcmut",
    "lc",
    "length",
    "editdistancetotrastuzumab",
}


def compute_release_metrics(
    metadata_sequences: Path,
    experimental: Path,
    literature: Path,
    affinity_archive: Path,
) -> dict[str, Any]:
    totals = {
        "total_antibody_entries": 0,
        "data_points": 0,
        "developability_annotations": 0,
        "source_records": 0,
    }
    entry_ids: set[str] = set()
    workbook_sequences: set[tuple[str, ...]] = set()
    affinity_sequences: set[tuple[str, ...]] = set()
    affinity_workbook_overlap: set[tuple[str, ...]] = set()
    components: list[dict[str, Any]] = []

    _process_workbook(
        metadata_sequences,
        {
            "metadata": "metadata",
            "sequences": "sequences",
        },
        totals,
        entry_ids,
        workbook_sequences,
        components,
    )
    _process_workbook(
        experimental,
        {
            name: "experimental_average"
            if name.startswith("Assay Data - average")
            else "experimental_raw"
            for name in _workbook_sheet_names(experimental)
            if name.startswith("Assay Data - average")
            or name.startswith("Assay Data - tidy format")
        },
        totals,
        entry_ids,
        workbook_sequences,
        components,
    )
    _process_workbook(
        literature,
        {
            name: "literature"
            for name in _workbook_sheet_names(literature)
            if name != "link"
        },
        totals,
        entry_ids,
        workbook_sequences,
        components,
    )

    for dataset_name, handle in iter_affinity_csv_streams(affinity_archive):
        try:
            component = _profile_affinity_stream(
                dataset_name,
                handle,
                workbook_sequences,
                affinity_sequences,
                affinity_workbook_overlap,
            )
        finally:
            handle.close()
        totals["total_antibody_entries"] += component["new_unique_entries"]
        totals["data_points"] += component["nonempty_cells"]
        totals["developability_annotations"] += component["annotation_count"]
        totals["source_records"] += component["row_count"]
        components.append(component)

    summary = {
        **totals,
        "dotad_workbook_entry_union": len(entry_ids),
        "dotad_workbook_sequence_fingerprints": len(workbook_sequences),
        "affinity_unique_sequence_fingerprints": len(affinity_sequences),
        "affinity_sequence_fingerprints_matching_workbooks": len(
            affinity_workbook_overlap
        ),
        "affinity_unique_entries_added_after_dedup": (
            totals["total_antibody_entries"] - len(entry_ids)
        ),
        "workbook_source_rows": sum(
            row["row_count"] for row in components if row["source_type"] == "workbook"
        ),
        "affinity_source_rows": sum(
            row["row_count"] for row in components if row["source_type"] == "affinity"
        ),
        "affinity_dataset_count": sum(
            row["source_type"] == "affinity" for row in components
        ),
        "sequence_unique_definition": (
            "Workbook records contribute unique antibody_id values. Affinity records "
            "are deduplicated by a normalized sequence fingerprint, preferring an "
            "ordered heavy/light pair, then heavy-only, single-sequence, CDR tuple, "
            "or ordered sequence tuple. Therefore the total is a hybrid identifier/"
            "sequence count, not a strict VH/VL-pair-unique count across all assets."
        ),
        "data_point_definition": (
            "Legacy-compatible heuristic: count every non-empty cell across scoped "
            "logical tables, including identifiers, sequences, metadata, and values. "
            "This is not a count of independent scientific observations."
        ),
        "developability_annotation_definition": (
            "Legacy-compatible heuristic: count non-empty fields classified as "
            "developability/affinity metrics. Raw, averaged, and source-specific "
            "representations are not deduplicated against one another."
        ),
        "metric_definition_status": "audit_reconstruction_not_release_definition",
    }
    return {"summary": summary, "components": components}


def iter_affinity_csv_streams(
    archive_path: Path,
) -> Iterator[tuple[str, TextIO]]:
    outer = zipfile.ZipFile(archive_path)
    try:
        for info in outer.infolist():
            normalized = info.filename.replace("\\", "/")
            if info.is_dir() or "__MACOSX/" in normalized or "/._" in normalized:
                continue
            lower = normalized.lower()
            if lower.endswith(".csv"):
                raw = outer.open(info)
                yield Path(normalized).name, io.TextIOWrapper(
                    raw,
                    encoding="utf-8-sig",
                    errors="replace",
                    newline="",
                )
            elif lower.endswith(".csv.zip"):
                payload = outer.read(info)
                nested = zipfile.ZipFile(io.BytesIO(payload))
                try:
                    members = [
                        member
                        for member in nested.infolist()
                        if not member.is_dir()
                        and member.filename.lower().endswith(".csv")
                        and "__MACOSX/" not in member.filename
                        and "/._" not in member.filename
                    ]
                    for member in members:
                        raw = nested.open(member)
                        wrapper = io.TextIOWrapper(
                            raw,
                            encoding="utf-8-sig",
                            errors="replace",
                            newline="",
                        )
                        yield Path(member.filename).name, wrapper
                finally:
                    nested.close()
    finally:
        outer.close()


def sequence_fingerprint(
    headers: Sequence[Any],
    row: Sequence[Any],
    sequence_indices: Sequence[int],
) -> tuple[str, ...] | None:
    values_by_header: dict[str, str] = {}
    ordered_values: list[str] = []
    for index in sequence_indices:
        if index >= len(row) or not looks_like_sequence(row[index]):
            continue
        value = normalize_sequence(row[index])
        compact = compact_header(headers[index])
        values_by_header.setdefault(compact, value)
        ordered_values.append(value)
    if not ordered_values:
        return None
    heavy = _first_sequence(values_by_header, HEAVY_SEQUENCE_HEADERS)
    light = _first_sequence(values_by_header, LIGHT_SEQUENCE_HEADERS)
    if heavy and light:
        return ("heavy_light", heavy, light)
    if heavy:
        return ("heavy_only", heavy)
    single = _first_sequence(values_by_header, SINGLE_SEQUENCE_HEADERS)
    if single:
        return ("sequence", single)
    cdr_values = [
        _first_sequence(values_by_header, group)
        for group in CDR_SEQUENCE_HEADER_GROUPS
    ]
    if any(cdr_values):
        return ("cdr", *cdr_values)
    return ("seq_tuple", *ordered_values)


def compact_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", normalize_text(value).lower())


def normalize_text(value: Any) -> str:
    return "" if value is None else str(value).replace("\ufeff", "").strip()


def normalize_sequence(value: Any) -> str:
    return normalize_text(value).upper().replace(" ", "").replace("-", "")


def looks_like_sequence(value: Any) -> bool:
    text = normalize_sequence(value)
    return (
        len(text) >= 6
        and not any(character.isdigit() for character in text)
        and all(character in AMINO_ACID_ALPHABET for character in text)
    )


def _process_workbook(
    path: Path,
    sheet_kinds: dict[str, str],
    totals: dict[str, int],
    entry_ids: set[str],
    workbook_sequences: set[tuple[str, ...]],
    components: list[dict[str, Any]],
) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        for sheet_name, kind in sheet_kinds.items():
            sheet = workbook[sheet_name]
            iterator = sheet.iter_rows(values_only=True)
            headers = next(iterator, ())
            id_index = next(
                (
                    index
                    for index, header in enumerate(headers)
                    if compact_header(header) in ENTRY_ID_HEADERS
                ),
                None,
            )
            sequence_indices = (
                [
                    index
                    for index, header in enumerate(headers)
                    if compact_header(header) in WORKBOOK_SEQUENCE_HEADERS
                ]
                if kind == "sequences"
                else []
            )
            annotation_indices = _workbook_annotation_indices(headers, kind)
            row_count = 0
            nonempty_cells = 0
            annotations = 0
            ids_in_sheet: set[str] = set()
            for row in iterator:
                if not any(normalize_text(value) for value in row):
                    continue
                row_count += 1
                nonempty_cells += sum(bool(normalize_text(value)) for value in row)
                if (
                    id_index is not None
                    and id_index < len(row)
                    and _is_meaningful(row[id_index])
                ):
                    ids_in_sheet.add(normalize_text(row[id_index]).upper())
                if sequence_indices:
                    fingerprint = sequence_fingerprint(
                        headers,
                        row,
                        sequence_indices,
                    )
                    if fingerprint:
                        workbook_sequences.add(fingerprint)
                annotations += sum(
                    index < len(row) and _is_meaningful(row[index])
                    for index in annotation_indices
                )
            new_ids = ids_in_sheet - entry_ids
            entry_ids.update(ids_in_sheet)
            totals["total_antibody_entries"] += len(new_ids)
            totals["data_points"] += nonempty_cells
            totals["developability_annotations"] += annotations
            totals["source_records"] += row_count
            components.append(
                {
                    "source_type": "workbook",
                    "file": path.name,
                    "table": sheet_name,
                    "row_count": row_count,
                    "nonempty_cells": nonempty_cells,
                    "annotation_count": annotations,
                    "unique_ids_in_table": len(ids_in_sheet),
                    "new_unique_entries": len(new_ids),
                    "unique_sequence_fingerprints": "",
                    "sequence_rows": "",
                    "metric_columns": len(annotation_indices),
                    "notes": f"Logical workbook table classified as {kind}.",
                }
            )
    finally:
        workbook.close()


def _profile_affinity_stream(
    dataset_name: str,
    handle: TextIO,
    workbook_sequences: set[tuple[str, ...]],
    affinity_sequences: set[tuple[str, ...]],
    affinity_workbook_overlap: set[tuple[str, ...]],
) -> dict[str, Any]:
    preview = [handle.readline() for _ in range(25)]
    preview = [line for line in preview if line != ""]
    header_offset = _detect_header_offset(preview)
    reader = csv.reader(chain(preview[header_offset:], handle))
    headers = next(reader, [])
    sequence_indices = _binding_sequence_indices(headers)
    metric_indices = _binding_metric_indices(headers)
    row_count = 0
    nonempty_cells = 0
    annotations = 0
    sequence_rows = 0
    file_sequences: set[tuple[str, ...]] = set()
    for row in reader:
        if not any(normalize_text(value) for value in row):
            continue
        row_count += 1
        nonempty_cells += sum(bool(normalize_text(value)) for value in row)
        fingerprint = sequence_fingerprint(headers, row, sequence_indices)
        if fingerprint:
            sequence_rows += 1
            file_sequences.add(fingerprint)
        annotations += sum(
            index < len(row) and _is_meaningful(row[index])
            for index in metric_indices
        )
    affinity_workbook_overlap.update(file_sequences & workbook_sequences)
    new_sequences = file_sequences - affinity_sequences - workbook_sequences
    affinity_sequences.update(file_sequences)
    return {
        "source_type": "affinity",
        "file": "dotad_affinity_benchmark_v2.0.zip",
        "table": dataset_name,
        "row_count": row_count,
        "nonempty_cells": nonempty_cells,
        "annotation_count": annotations,
        "unique_ids_in_table": "",
        "new_unique_entries": len(new_sequences),
        "unique_sequence_fingerprints": len(file_sequences),
        "sequence_rows": sequence_rows,
        "metric_columns": len(metric_indices),
        "notes": (
            f"Header line offset={header_offset}; sequence fingerprints are "
            "audit-normalized but source values are unchanged."
        ),
    }


def _workbook_sheet_names(path: Path) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _workbook_annotation_indices(
    headers: Sequence[Any],
    kind: str,
) -> list[int]:
    if kind == "metadata":
        return [
            index
            for index, header in enumerate(headers)
            if compact_header(header) == "adasrate"
        ]
    if kind == "sequences":
        return []
    if kind == "experimental_average":
        return [
            index
            for index, header in enumerate(headers)
            if compact_header(header) not in EXPERIMENTAL_META_HEADERS
            and not any(
                pattern in normalize_text(header).lower()
                for pattern in EXPERIMENTAL_SUPPORT_PATTERNS
            )
            and "avg" in normalize_text(header).lower()
        ]
    if kind == "experimental_raw":
        return [
            index
            for index, header in enumerate(headers)
            if compact_header(header) not in EXPERIMENTAL_META_HEADERS
            and not any(
                pattern in normalize_text(header).lower()
                for pattern in EXPERIMENTAL_SUPPORT_PATTERNS
            )
        ]
    if kind == "literature":
        return [
            index
            for index, header in enumerate(headers)
            if compact_header(header)
            and compact_header(header) not in LITERATURE_META_HEADERS
        ]
    return []


def _binding_sequence_indices(headers: Sequence[Any]) -> list[int]:
    output: list[int] = []
    for index, header in enumerate(headers):
        raw = normalize_text(header).lower()
        compact = compact_header(header)
        if "antigen" in raw or compact.startswith("ag"):
            continue
        if (
            compact in BINDING_SEQUENCE_HEADERS
            or "sequence" in raw
            or compact.startswith(("hcdr", "lcdr", "cdrh", "cdrl"))
        ):
            output.append(index)
    return output


def _binding_metric_indices(headers: Sequence[Any]) -> list[int]:
    output: list[int] = []
    for index, header in enumerate(headers):
        raw = normalize_text(header).lower()
        compact = compact_header(header)
        if not compact or compact in BINDING_META_HEADERS:
            continue
        if compact in BINDING_SEQUENCE_HEADERS:
            continue
        if "antigen" in raw or compact.startswith("ag") or "sequence" in raw:
            continue
        if compact.startswith(("hcdr", "lcdr", "cdrh", "cdrl")):
            continue
        if any(pattern in raw for pattern in BINDING_SUPPORT_PATTERNS):
            continue
        if any(keyword in raw for keyword in BINDING_METRIC_KEYWORDS):
            output.append(index)
    return output


def _detect_header_offset(lines: list[str]) -> int:
    best_index = 0
    best_score = -1
    for index, line in enumerate(lines):
        if not line:
            continue
        row = next(csv.reader([line]))
        score = int(len(row) >= 2)
        score += 2 * int(
            any(compact_header(value) in BINDING_SEQUENCE_HEADERS for value in row)
        )
        lowered = [normalize_text(value).lower() for value in row]
        score += 2 * int(
            any(
                any(keyword in value for keyword in BINDING_METRIC_KEYWORDS)
                for value in lowered
            )
        )
        score += int(
            any(compact_header(value) in {"predaffinity", "fitness"} for value in row)
        )
        if score > best_score:
            best_score = score
            best_index = index
    return best_index


def _first_sequence(
    values: dict[str, str],
    candidates: Iterable[str],
) -> str:
    return next((values[name] for name in candidates if values.get(name)), "")


def _is_meaningful(value: Any) -> bool:
    return normalize_text(value).lower() not in PLACEHOLDER_TOKENS
