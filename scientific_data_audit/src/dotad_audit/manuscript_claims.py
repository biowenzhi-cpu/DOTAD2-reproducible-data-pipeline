from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import yaml
from docx import Document


NUMBER_PATTERN = re.compile(
    r"(?<![\w.])(?:approximately\s+|about\s+|~\s*)?"
    r"(?:\d{1,3}(?:,\d{3})+|\d+(?:\.\d+)?)"
    r"(?:\s*(?:million|billion|%|fold|pairs?|entries|records?|antibodies|assays?))?",
    re.IGNORECASE,
)
REFERENCE_HEADING = re.compile(r"^\s*(references|bibliography)\s*$", re.IGNORECASE)
SECTION_PATTERN = re.compile(r"^\s*(?:\d+(?:\.\d+)*\s+.+|Abstract|Introduction|Methods|Results|Discussion|Supplementary materials)\s*$", re.IGNORECASE)


def extract_document_paragraphs(path: Path) -> list[dict[str, Any]]:
    document = Document(path)
    output: list[dict[str, Any]] = []
    section = "Front matter"
    in_references = False
    paragraph_index = 0
    for paragraph in document.paragraphs:
        paragraph_index += 1
        text = paragraph.text.strip()
        if not text:
            continue
        if REFERENCE_HEADING.match(text):
            section = "References"
            in_references = True
        elif SECTION_PATTERN.match(text) and len(text) < 180:
            section = text
            in_references = section.lower() in {"references", "bibliography"}
        output.append(
            {
                "paragraph_index": paragraph_index,
                "section": section,
                "text": text,
                "is_reference_section": in_references,
            }
        )
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if text:
                paragraph_index += 1
                output.append(
                    {
                        "paragraph_index": paragraph_index,
                        "section": f"Table {table_index}",
                        "text": text,
                        "is_reference_section": False,
                        "location_note": f"table {table_index}, row {row_index}",
                    }
                )
    return output


def extract_numeric_claims(
    manuscript_path: Path,
    registry_path: Path,
) -> list[dict[str, Any]]:
    paragraphs = extract_document_paragraphs(manuscript_path)
    registry = (yaml.safe_load(registry_path.read_text(encoding="utf-8")) or {}).get("claims", [])
    claims: list[dict[str, Any]] = []
    matched_spans: dict[int, list[tuple[int, int]]] = {}
    next_id = 1
    for registered in registry:
        found = False
        for paragraph in paragraphs:
            if paragraph["is_reference_section"]:
                continue
            text = paragraph["text"]
            for pattern in registered.get("patterns", []):
                match = re.search(pattern, text, re.IGNORECASE)
                if not match:
                    continue
                claims.append(
                    _claim_row(
                        claim_id=str(registered["id"]),
                        paragraph=paragraph,
                        claim_text=_context(text, match.start(), match.end()),
                        metric_name=str(registered.get("metric", registered["id"])),
                        claimed_value=registered.get("value", ""),
                        claimed_unit=str(registered.get("unit", "")),
                        category=str(registered.get("category", "other")),
                        priority=str(registered.get("priority", "medium")),
                        notes="Registered priority claim.",
                    )
                )
                matched_spans.setdefault(paragraph["paragraph_index"], []).append(match.span())
                found = True
                break
            if found:
                break
        if not found:
            claims.append(
                {
                    "claim_id": str(registered["id"]),
                    "section": "Not located in extractable DOCX body text",
                    "paragraph_index": "",
                    "page_or_location": "Not located; may be figure-only, text-box content, or absent",
                    "claim_text": (
                        "Registered priority claim was not found in extractable manuscript body/table text."
                    ),
                    "metric_name": str(registered.get("metric", registered["id"])),
                    "claimed_value": registered.get("value", ""),
                    "claimed_unit": str(registered.get("unit", "")),
                    "claim_category": str(registered.get("category", "other")),
                    "priority": str(registered.get("priority", "medium")),
                    "requires_data_reproduction": True,
                    "notes": (
                        "Retained in the audit registry because the requested audit explicitly prioritizes "
                        "this value; absence from extractable text is not evidence of correctness."
                    ),
                }
            )

    for paragraph in paragraphs:
        if paragraph["is_reference_section"]:
            continue
        text = paragraph["text"]
        if paragraph["section"] == "Front matter":
            continue
        if SECTION_PATTERN.match(text) and len(text) < 180:
            continue
        for match in NUMBER_PATTERN.finditer(text):
            if _is_excluded_number(text, match):
                continue
            if any(_overlap(match.span(), span) for span in matched_spans.get(paragraph["paragraph_index"], [])):
                continue
            rendered = match.group(0).strip()
            value, unit = parse_numeric_token(rendered)
            claims.append(
                _claim_row(
                    claim_id=f"auto_{next_id:04d}",
                    paragraph=paragraph,
                    claim_text=_context(text, match.start(), match.end()),
                    metric_name="automatically extracted numeric statement",
                    claimed_value=value,
                    claimed_unit=unit,
                    category=infer_claim_category(text),
                    priority="medium",
                    notes="Automatically extracted; requires manual relevance review.",
                )
            )
            next_id += 1
    return claims


def parse_numeric_token(token: str) -> tuple[float | int | str, str]:
    lowered = token.lower().replace(",", "").strip()
    unit = ""
    for suffix in ("million", "billion", "%", "fold", "pairs", "pair", "entries", "records", "record", "antibodies", "assays"):
        if suffix in lowered:
            unit = suffix
            lowered = lowered.replace(suffix, "").strip()
            break
    lowered = re.sub(r"^(approximately|about|~)\s*", "", lowered).strip()
    try:
        number = float(lowered)
        if unit == "million":
            number *= 1_000_000
        elif unit == "billion":
            number *= 1_000_000_000
        return (int(number) if number.is_integer() else number), unit
    except ValueError:
        return token, unit


def infer_claim_category(text: str) -> str:
    lowered = text.lower()
    if any(token in lowered for token in ("pearson", "spearman", "correlation", "r =")):
        return "correlation"
    if any(token in lowered for token in ("pca", "pc1", "pc2", "pc3", "variance")):
        return "pca"
    if any(token in lowered for token in ("p <", "p =", "wilcoxon", "test")):
        return "statistical_test"
    if "concordance" in lowered:
        return "concordance"
    if any(token in lowered for token in ("coverage", "covered")):
        return "coverage"
    if any(token in lowered for token in ("missing", "completeness")):
        return "missingness"
    if any(token in lowered for token in ("benchmark", "task")):
        return "benchmark"
    if any(token in lowered for token in ("entry", "record", "annotation", "resource")):
        return "resource_size"
    if "%" in text or "format" in lowered:
        return "composition"
    return "other"


def claim_context_markdown(claims: list[dict[str, Any]], manuscript_path: Path) -> str:
    lines = [
        "# Manuscript Numeric Claim Context",
        "",
        f"- Manuscript: `{manuscript_path}`",
        f"- Extracted claims: {len(claims)}",
        "- Reference-list numbers, citation-only brackets, software versions, and author-affiliation markers are excluded heuristically.",
        "- Automatically extracted claims require manual relevance review; registered claims are prioritized for mapping.",
        "",
    ]
    for claim in claims:
        claimed_value_line = f"- Claimed value: `{claim['claimed_value']}`"
        if claim["claimed_unit"]:
            claimed_value_line += f" {claim['claimed_unit']}"
        lines.extend(
            [
                f"## {claim['claim_id']} — {claim['metric_name']}",
                "",
                f"- Section: {claim['section']}",
                f"- Paragraph index: {claim['paragraph_index'] or 'not located'}",
                claimed_value_line,
                f"- Category / priority: `{claim['claim_category']}` / `{claim['priority']}`",
                f"- Context: {str(claim['claim_text']).strip()}",
                "",
            ]
        )
    return "\n".join(lines)


def _claim_row(
    *,
    claim_id: str,
    paragraph: dict[str, Any],
    claim_text: str,
    metric_name: str,
    claimed_value: Any,
    claimed_unit: str,
    category: str,
    priority: str,
    notes: str,
) -> dict[str, Any]:
    return {
        "claim_id": claim_id,
        "section": paragraph["section"],
        "paragraph_index": paragraph["paragraph_index"],
        "page_or_location": paragraph.get("location_note", "DOCX paragraph order; page unavailable without rendering"),
        "claim_text": claim_text,
        "metric_name": metric_name,
        "claimed_value": claimed_value,
        "claimed_unit": claimed_unit,
        "claim_category": category,
        "priority": priority,
        "requires_data_reproduction": True,
        "notes": notes,
    }


def _context(text: str, start: int, end: int, radius: int = 180) -> str:
    left = max(0, start - radius)
    right = min(len(text), end + radius)
    return text[left:right]


def _is_excluded_number(text: str, match: re.Match[str]) -> bool:
    token = match.group(0).strip()
    start, end = match.span()
    before = text[max(0, start - 40) : start]
    after = text[end : min(len(text), end + 40)]
    bracket_start = text.rfind("[", max(0, start - 30), start + 1)
    bracket_end = text.find("]", end, min(len(text), end + 30))
    if bracket_start != -1 and bracket_end != -1:
        bracket_content = text[bracket_start + 1 : bracket_end]
        if re.fullmatch(r"[\d,\s;–-]+", bracket_content):
            return True
    if re.fullmatch(r"\d{4}", re.sub(r"\D", "", token)) and re.search(r"(19|20)\d{2}", token):
        return True
    if re.search(r"\[[\d,\s–-]+\]", text[max(0, start - 2) : min(len(text), end + 2)]):
        return True
    if re.search(r"(?i)(version|v\.?|python|r)\s*$", before) and re.match(r"\d+(?:\.\d+)+", token):
        return True
    if re.search(r"(?i)DOTAD[-\s]*$", before) and re.match(r"\d+(?:\.\d+)*", token):
        return True
    if re.match(r"^\.\d", token) or re.search(r"\d\.$", before):
        return True
    if re.search(r"(?i)(fig(?:ure)?s?|tables?|sections?)\s*$", before) and re.fullmatch(r"\d+(?:\.\d+)?", token):
        return True
    if re.search(r"(?i)figures?\s+\d+\s+(?:and|to|[-–])\s*$", before) and re.fullmatch(r"\d+", token):
        return True
    if re.search(r"(?i)(pmid|doi|grant|fund|project|award|no\.)\s*[:#]?\s*$", before):
        return True
    local_context = text[max(0, start - 80) : min(len(text), end + 80)]
    if re.search(r"https?://\S*$", text[max(0, start - 120) : start]):
        return True
    if re.search(r"(?i)\b(PMID|DOI)\b", local_context):
        return True
    if re.search(r"(?i)(grant\s+nos?\.?|funded by|funding|natural science foundation)", local_context):
        return True
    if re.search(
        r"(?i)(January|February|March|April|May|June|July|August|September|October|November|December)"
        r"\s+\d{1,2},\s+\d{4}",
        local_context,
    ):
        return True
    if before.endswith("(") and after.startswith(")") and re.fullmatch(r"\d{1,3}", token):
        return True
    return False


def _overlap(first: tuple[int, int], second: tuple[int, int]) -> bool:
    return first[0] < second[1] and second[0] < first[1]
