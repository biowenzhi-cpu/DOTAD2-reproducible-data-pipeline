from __future__ import annotations

import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def profile_workbook(
    path: Path,
    file_id: str,
    *,
    max_samples: int = 3,
    max_exact_unique: int = 100_000,
) -> tuple[dict[str, Any], list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    workbook_row = {
        "file_id": file_id,
        "relative_path": "",
        "workbook_type": path.suffix.lower(),
        "sheet_count": "",
        "has_vba": path.suffix.lower() == ".xlsm",
        "read_status": "unreadable",
        "notes": "",
    }
    sheet_rows: list[dict[str, Any]] = []
    table_rows: list[dict[str, Any]] = []
    column_rows: list[dict[str, Any]] = []
    try:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=False,
            keep_vba=False,
            keep_links=False,
        )
        cached_workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
            keep_vba=False,
            keep_links=False,
        )
        workbook_row["sheet_count"] = len(workbook.sheetnames)
        workbook_row["read_status"] = "readable"
        for sheet_index, sheet_name in enumerate(workbook.sheetnames, start=1):
            sheet = workbook[sheet_name]
            cached_sheet = cached_workbook[sheet_name]
            sheet_row, table_row, columns = profile_sheet(
                sheet,
                cached_sheet,
                f"{file_id}:S{sheet_index:03d}",
                max_samples=max_samples,
                max_exact_unique=max_exact_unique,
            )
            sheet_row.update(
                {
                    "file_id": file_id,
                    "sheet_name": sheet_name,
                    "sheet_index": sheet_index,
                    "sheet_visibility": sheet.sheet_state,
                }
            )
            table_row["sheet_or_table_name"] = sheet_name
            sheet_rows.append(sheet_row)
            table_rows.append(table_row)
            column_rows.extend(columns)
        workbook.close()
        cached_workbook.close()
    except Exception as exc:  # openpyxl raises several format/corruption types
        workbook_row["notes"] = f"{type(exc).__name__}: {exc}"
    return workbook_row, sheet_rows, table_rows, column_rows


def profile_sheet(
    sheet: Any,
    cached_sheet: Any,
    table_id: str,
    *,
    max_samples: int = 3,
    max_exact_unique: int = 100_000,
) -> tuple[dict[str, Any], dict[str, Any], list[dict[str, Any]]]:
    physical_max_row = int(sheet.max_row or 0)
    physical_max_column = int(sheet.max_column or 0)
    nonempty_rows: list[tuple[int, list[Any]]] = []
    formula_count = 0
    cached_count = 0
    blank_rows = 0
    for row_index, row in enumerate(sheet.iter_rows(), start=1):
        values = [cell.value for cell in row]
        has_value = any(value is not None and str(value).strip() != "" for value in values)
        if has_value:
            nonempty_rows.append((row_index, values))
        else:
            blank_rows += 1
        for column_index, cell in enumerate(row, start=1):
            if cell.data_type == "f" or (
                isinstance(cell.value, str) and cell.value.startswith("=")
            ):
                formula_count += 1
                cached_value = cached_sheet.cell(row=row_index, column=column_index).value
                if cached_value is not None:
                    cached_count += 1

    if nonempty_rows:
        header_row_number, header_values = nonempty_rows[0]
        data_rows = [values for row_number, values in nonempty_rows if row_number > header_row_number]
        header_confidence = _header_confidence(header_values)
    else:
        header_row_number, header_values, data_rows, header_confidence = "", [], [], 0.0
    last_nonempty_column = max(
        (
            index + 1
            for _, values in nonempty_rows
            for index, value in enumerate(values)
            if value is not None and str(value).strip() != ""
        ),
        default=0,
    )
    columns = [
        str(value).strip() if value is not None and str(value).strip() else f"unnamed_{index + 1}"
        for index, value in enumerate(header_values[:last_nonempty_column])
    ]
    row_tuples: list[tuple[str, ...]] = []
    stats = [_new_stat(column) for column in columns]
    fully_empty_rows = 0
    for values in data_rows:
        normalized = tuple(
            _render(values[index] if index < len(values) else None)
            for index in range(len(columns))
        )
        row_tuples.append(normalized)
        if not any(normalized):
            fully_empty_rows += 1
        for index, value in enumerate(normalized):
            _update_stat(stats[index], value, max_samples, max_exact_unique)
    duplicate_count = len(row_tuples) - len(set(row_tuples))
    roles = infer_candidate_columns(columns)
    sheet_row = {
        "table_id": table_id,
        "physical_max_row": physical_max_row,
        "physical_max_column": physical_max_column,
        "estimated_data_rows": len(data_rows),
        "estimated_data_columns": len(columns),
        "header_row": header_row_number,
        "header_confidence": header_confidence,
        "merged_range_count": len(getattr(sheet, "merged_cells", []).ranges)
        if hasattr(getattr(sheet, "merged_cells", None), "ranges")
        else 0,
        "formula_cell_count": formula_count,
        "cached_formula_value_count": cached_count,
        "blank_row_count": blank_rows,
        "read_status": "readable",
        "notes": (
            "Estimated rows count only non-empty worksheet rows after the detected header; "
            "formatted-but-empty rows are excluded. Formula caches are read from data_only mode."
        ),
    }
    table_row = {
        "table_id": table_id,
        "row_count": len(data_rows),
        "column_count": len(columns),
        "column_names_json": json.dumps(columns, ensure_ascii=False),
        "duplicate_full_row_count": duplicate_count,
        "fully_empty_row_count": fully_empty_rows,
        "fully_empty_column_count": sum(stat["non_missing_count"] == 0 for stat in stats),
        **{key: json.dumps(value, ensure_ascii=False) for key, value in roles.items()},
        "read_status": "readable",
        "notes": "Exact profile of non-empty worksheet region using openpyxl read_only mode.",
    }
    column_rows = [
        {
            "table_id": table_id,
            "column_name": stat["column_name"],
            "normalized_column_name": normalize_column_name(stat["column_name"]),
            "inferred_data_type": _dominant_type(stat["types"]),
            "non_missing_count": stat["non_missing_count"],
            "missing_count": max(0, len(data_rows) - stat["non_missing_count"]),
            "missing_fraction": (
                round((len(data_rows) - stat["non_missing_count"]) / len(data_rows), 6)
                if data_rows
                else 0.0
            ),
            "unique_count": len(stat["unique"]),
            "sample_values_json": json.dumps(
                _safe_samples(stat["samples"], infer_semantic_role(stat["column_name"])),
                ensure_ascii=False,
            ),
            "semantic_role_guess": infer_semantic_role(stat["column_name"]),
            "confidence": role_confidence(stat["column_name"]),
            "notes": "unique_count is a lower bound"
            if stat["unique_capped"]
            else "",
        }
        for stat in stats
    ]
    return sheet_row, table_row, column_rows


def read_sheet_records(path: Path, sheet_name: str, row_limit: int | None = None) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    header: list[str] = []
    for row in iterator:
        if any(value is not None and str(value).strip() for value in row):
            header = [
                str(value).strip() if value is not None and str(value).strip() else f"unnamed_{index + 1}"
                for index, value in enumerate(row)
            ]
            break
    records: list[dict[str, Any]] = []
    for row in iterator:
        if not any(value is not None and str(value).strip() for value in row):
            continue
        records.append({header[index]: value for index, value in enumerate(row[: len(header)])})
        if row_limit is not None and len(records) >= row_limit:
            break
    workbook.close()
    return records


def infer_candidate_columns(columns: list[str]) -> dict[str, list[str]]:
    mapping = {
        "candidate_primary_keys_json": [],
        "candidate_foreign_keys_json": [],
        "candidate_source_columns_json": [],
        "candidate_doi_columns_json": [],
        "candidate_pmid_columns_json": [],
        "candidate_dotad_id_columns_json": [],
        "candidate_sequence_columns_json": [],
        "candidate_assay_columns_json": [],
        "candidate_value_columns_json": [],
        "candidate_unit_columns_json": [],
        "candidate_condition_columns_json": [],
        "candidate_version_columns_json": [],
    }
    for column in columns:
        normalized = normalize_column_name(column)
        if normalized == "id" or normalized.endswith("_id"):
            mapping["candidate_primary_keys_json"].append(column)
            mapping["candidate_foreign_keys_json"].append(column)
        if "source" in normalized or "reference" in normalized or "study" in normalized:
            mapping["candidate_source_columns_json"].append(column)
        if "doi" in normalized:
            mapping["candidate_doi_columns_json"].append(column)
        if "pmid" in normalized or "pubmed" in normalized:
            mapping["candidate_pmid_columns_json"].append(column)
        if "dotad" in normalized and "id" in normalized or normalized == "antibody_id":
            mapping["candidate_dotad_id_columns_json"].append(column)
        if any(token in normalized for token in ("sequence", "_vh", "_vl", "_hc", "_lc", "cdr")):
            mapping["candidate_sequence_columns_json"].append(column)
        if any(token in normalized for token in ("assay", "endpoint", "metric")):
            mapping["candidate_assay_columns_json"].append(column)
        if any(token in normalized for token in ("value", "score", "measurement", "mean", "median")):
            mapping["candidate_value_columns_json"].append(column)
        if "unit" in normalized:
            mapping["candidate_unit_columns_json"].append(column)
        if any(token in normalized for token in ("condition", "ph", "temperature", "concentration")):
            mapping["candidate_condition_columns_json"].append(column)
        if "version" in normalized or "release" in normalized:
            mapping["candidate_version_columns_json"].append(column)
    return mapping


def normalize_column_name(value: str) -> str:
    import re

    normalized = re.sub(r"[^a-z0-9]+", "_", str(value).strip().lower()).strip("_")
    return normalized


def infer_semantic_role(column: str) -> str:
    roles = infer_candidate_columns([column])
    candidates = [key.removeprefix("candidate_").removesuffix("_columns_json") for key, values in roles.items() if values]
    return "|".join(candidates) if candidates else "unknown"


def role_confidence(column: str) -> str:
    normalized = normalize_column_name(column)
    if normalized in {"dotad_id", "antibody_id", "pmid", "doi", "unit", "value", "sequence"}:
        return "high"
    return "medium" if infer_semantic_role(column) != "unknown" else "low"


def _new_stat(column: str) -> dict[str, Any]:
    return {
        "column_name": column,
        "non_missing_count": 0,
        "unique": set(),
        "unique_capped": False,
        "samples": [],
        "types": Counter(),
    }


def _update_stat(stat: dict[str, Any], value: str, max_samples: int, max_unique: int) -> None:
    if value == "":
        return
    stat["non_missing_count"] += 1
    if len(stat["samples"]) < max_samples and value not in stat["samples"]:
        stat["samples"].append(value[:160])
    if len(stat["unique"]) < max_unique:
        stat["unique"].add(value)
    else:
        stat["unique_capped"] = True
    stat["types"][_infer_type(value)] += 1


def _render(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _infer_type(value: str) -> str:
    if value.lower() in {"true", "false"}:
        return "boolean"
    try:
        int(value)
        return "integer"
    except ValueError:
        pass
    try:
        float(value)
        return "number"
    except ValueError:
        return "string"


def _dominant_type(counter: Counter[str]) -> str:
    return counter.most_common(1)[0][0] if counter else "empty"


def _header_confidence(values: list[Any]) -> float:
    nonempty = [value for value in values if value is not None and str(value).strip()]
    if not nonempty:
        return 0.0
    strings = sum(isinstance(value, str) for value in nonempty)
    unique = len({str(value).strip().lower() for value in nonempty})
    return round(0.5 * strings / len(nonempty) + 0.5 * unique / len(nonempty), 3)


def _safe_samples(samples: list[str], semantic_role: str) -> list[str]:
    if "sequence" in semantic_role:
        return ["[sequence sample redacted from audit dictionary]"] if samples else []
    output = []
    for sample in samples:
        compact = sample.replace("\r", " ").replace("\n", " ").replace("\t", " ")
        output.append(compact[:160])
    return output
