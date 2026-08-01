from docx import Document

from dotad_audit.manuscript_claims import extract_numeric_claims


def _registry(path):
    path.write_text(
        """claims:
  - id: resource_entries
    metric: resource entries
    patterns: ["324,000\\\\s+sequence-unique"]
    value: 324000
    unit: entries
    category: resource_size
    priority: critical
""",
        encoding="utf-8",
    )


def test_docx_numeric_claim_extraction(tmp_path):
    manuscript = tmp_path / "manuscript.docx"
    registry = tmp_path / "claims.yml"
    _registry(registry)
    document = Document()
    document.add_paragraph("Results")
    document.add_paragraph(
        "DOTAD 2.0 contains approximately 324,000 sequence-unique antibody entries."
    )
    document.add_paragraph("Coverage reached 68% in the selected subset.")
    document.save(manuscript)
    claims = extract_numeric_claims(manuscript, registry)
    assert any(claim["claim_id"] == "resource_entries" for claim in claims)
    assert any(claim["claimed_value"] == 68 and claim["claimed_unit"] == "%" for claim in claims)


def test_reference_numbers_are_excluded(tmp_path):
    manuscript = tmp_path / "references.docx"
    registry = tmp_path / "claims.yml"
    _registry(registry)
    document = Document()
    document.add_paragraph("Introduction")
    document.add_paragraph("A claim supported by prior work [12].")
    document.add_paragraph("References")
    document.add_paragraph("Smith J. Journal 2024; 12: 100-110.")
    document.save(manuscript)
    claims = extract_numeric_claims(manuscript, registry)
    assert not any(claim["section"] == "References" for claim in claims)
    assert not any(claim["claimed_value"] == 12 for claim in claims)
