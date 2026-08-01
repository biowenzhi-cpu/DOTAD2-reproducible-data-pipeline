from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from dotad_audit.spreadsheet_profiler import profile_sheet


def _build_fixture(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "data"
    sheet.append(["id", "value"])
    sheet.append(["A", 2])
    sheet.append(["B", "=B2*2"])
    sheet["A10"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    workbook.save(path)


def test_formatted_blank_rows_are_not_counted(tmp_path):
    path = tmp_path / "formatted.xlsx"
    _build_fixture(path)
    workbook = load_workbook(path, read_only=True, data_only=False)
    cached = load_workbook(path, read_only=True, data_only=True)
    sheet_row, table_row, _ = profile_sheet(workbook["data"], cached["data"], "T1")
    workbook.close()
    cached.close()
    assert sheet_row["physical_max_row"] == 10
    assert sheet_row["estimated_data_rows"] == 2
    assert table_row["row_count"] == 2


def test_formula_cells_are_counted(tmp_path):
    path = tmp_path / "formula.xlsx"
    _build_fixture(path)
    workbook = load_workbook(path, read_only=True, data_only=False)
    cached = load_workbook(path, read_only=True, data_only=True)
    sheet_row, _, _ = profile_sheet(workbook["data"], cached["data"], "T1")
    workbook.close()
    cached.close()
    assert sheet_row["formula_cell_count"] == 1
    assert sheet_row["cached_formula_value_count"] == 0
