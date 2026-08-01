from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path
from typing import Any, Mapping, Sequence

import openpyxl

from .workbook_reader import serialize_scalar


LICENCE_STATUSES = {
    "EXPLICIT_FILE_LEVEL_LICENCE_FOUND",
    "PACKAGE_LEVEL_LICENCE_FOUND",
    "PUBLICATION_ONLY_LICENCE_FOUND",
    "NO_LOCAL_LICENCE_EVIDENCE",
    "CONFLICTING_LICENCE_EVIDENCE",
}

RELEASE_STATUSES = {
    "LITERATURE_COMPILATION_WITH_ATTRIBUTION",
    "LINK_ONLY",
    "EXCLUDE_COPYRIGHTED_CONTENT",
    "NEEDS_MANUAL_REVIEW",
}


def audit_gdpa_files(
    candidate_files: Mapping[str, Path],
    publication_licences: Mapping[str, str] | None = None,
) -> list[dict[str, Any]]:
    publication_licences = publication_licences or {}
    rows = []
    for dataset_id, candidate in sorted(candidate_files.items()):
        path = Path(candidate)
        evidence_location = ""
        excerpt = ""
        licence_name = ""
        status = "NO_LOCAL_LICENCE_EVIDENCE"
        notes = "No explicit file- or package-level licence text was found locally."
        inspected = ""
        if path.exists() and path.suffix.casefold() in {".xlsx", ".xlsm"}:
            (
                evidence_location,
                excerpt,
                licence_name,
                inspected,
            ) = _workbook_licence(path)
        if excerpt:
            status = (
                "CONFLICTING_LICENCE_EVIDENCE"
                if licence_name == "CONFLICTING"
                else "EXPLICIT_FILE_LEVEL_LICENCE_FOUND"
            )
            notes = (
                "Conflicting licence wording was located in the candidate file."
                if status == "CONFLICTING_LICENCE_EVIDENCE"
                else "Explicit licence wording was located in the candidate file."
            )
        elif publication_licences.get(dataset_id):
            status = "PUBLICATION_ONLY_LICENCE_FOUND"
            notes = (
                "A publication licence was recorded separately and was not "
                "propagated to the dataset file."
            )
        if not excerpt and inspected:
            evidence_location = inspected
            notes += f" Inspected: {inspected}."
        rows.append(
            {
                "dataset_id": dataset_id,
                "candidate_file": path.name,
                "evidence_location": evidence_location,
                "licence_name_as_written": licence_name,
                "licence_text_excerpt": excerpt,
                "applies_to_file": "true" if excerpt else "false",
                "redistribution_statement": "",
                "derivative_use_statement": "",
                "commercial_use_statement": "",
                "attribution_statement": "",
                "evidence_status": status,
                "notes": notes,
            }
        )
    return rows


def build_literature_release_status(
    literature_rows: Sequence[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    by_sheet: dict[str, list[Mapping[str, Any]]] = defaultdict(list)
    for row in literature_rows:
        by_sheet[serialize_scalar(row.get("source_sheet")).strip()].append(row)
    output = []
    for sheet, rows in sorted(by_sheet.items()):
        has_long_text = any(_has_long_text(row) for row in rows)
        citation_available = all(
            any(
                serialize_scalar(row.get(field)).strip()
                for field in ("pmid", "doi", "source_url")
            )
            for row in rows
        )
        if has_long_text:
            status = "EXCLUDE_COPYRIGHTED_CONTENT"
        elif citation_available:
            status = "LITERATURE_COMPILATION_WITH_ATTRIBUTION"
        else:
            status = "NEEDS_MANUAL_REVIEW"
        output.append(
            {
                "source_id": serialize_scalar(rows[0].get("source_id")),
                "source_sheet": sheet,
                "record_count": len(rows),
                "contains_factual_values": "true",
                "contains_copied_long_text": "true" if has_long_text else "false",
                "contains_copied_figures": "false",
                "contains_copied_table_layout": "false",
                "row_level_source_citation_available": (
                    "true" if citation_available else "false"
                ),
                "release_status": status,
                "attribution_required": "true",
                "manual_review_required": (
                    "true" if status == "NEEDS_MANUAL_REVIEW" else "false"
                ),
                "notes": (
                    "Status concerns the DOTAD factual compilation only; it "
                    "does not grant rights to redistribute source articles, "
                    "figures, or original table layouts."
                ),
            }
        )
    return output


def _workbook_licence(path: Path) -> tuple[str, str, str, str]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    matches = []
    inspected = []
    try:
        properties = workbook.properties
        property_values = {
            "title": properties.title,
            "subject": properties.subject,
            "description": properties.description,
            "keywords": properties.keywords,
            "category": properties.category,
        }
        inspected.append("workbook properties")
        for field, value in property_values.items():
            text = serialize_scalar(value).strip()
            if text and re.search(
                r"\blicen[cs](?:e|ed)\b|creative commons|\bcc[- ]by\b",
                text,
                flags=re.IGNORECASE,
            ):
                matches.append((f"properties.{field}", 0, text))
        for sheet_name in workbook.sheetnames:
            if re.search(
                r"definition|version|licen[cs]e|terms|readme|citation|"
                r"dataset.?card|manifest",
                sheet_name,
                flags=re.IGNORECASE,
            ):
                inspected.append(f"sheet {sheet_name}")
            if not re.search(
                r"licen[cs]e|terms|readme|citation|dataset.?card",
                sheet_name,
                flags=re.IGNORECASE,
            ):
                continue
            sheet = workbook[sheet_name]
            for row_index, row in enumerate(
                sheet.iter_rows(values_only=True),
                start=1,
            ):
                for value in row:
                    text = serialize_scalar(value).strip()
                    if re.search(
                        r"\blicen[cs](?:e|ed)\b|creative commons|\bcc[- ]by\b",
                        text,
                        flags=re.IGNORECASE,
                    ):
                        matches.append((sheet_name, row_index, text))
        if not matches:
            return "", "", "", "; ".join(inspected)
        locations = {
            f"{sheet}!row {row}" for sheet, row, _ in matches
        }
        texts = {text for _, _, text in matches}
        if len(texts) > 1 and _conflicting_texts(texts):
            excerpt = " | ".join(sorted(texts))[:300]
            return (
                "; ".join(sorted(locations)),
                excerpt,
                "CONFLICTING",
                "; ".join(inspected),
            )
        text = matches[0][2]
        excerpt = text[:300]
        licence = _licence_name(text)
        return (
            f"{path.name}::{matches[0][0]}!row {matches[0][1]}",
            excerpt,
            licence,
            "; ".join(inspected),
        )
    finally:
        workbook.close()


def _licence_name(text: str) -> str:
    match = re.search(
        r"(CC[- ]BY(?:[- ]NC)?(?:[- ]SA)?(?:\s*\d(?:\.\d)?)?)",
        text,
        flags=re.IGNORECASE,
    )
    return match.group(1).upper().replace("-", " ") if match else "AS_WRITTEN"


def _conflicting_texts(texts: set[str]) -> bool:
    normalized = " ".join(texts).casefold()
    return "all rights reserved" in normalized and (
        "cc by" in normalized or "creative commons" in normalized
    )


def _has_long_text(row: Mapping[str, Any]) -> bool:
    payload_text = serialize_scalar(row.get("source_payload_json"))
    try:
        payload = json.loads(payload_text)
    except (json.JSONDecodeError, TypeError):
        payload = []
    for item in payload if isinstance(payload, list) else []:
        value = serialize_scalar(item.get("value"))
        if len(value) > 1000 or len(value.split()) > 100:
            return True
    return False
