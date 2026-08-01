from __future__ import annotations

import json
from pathlib import Path

import openpyxl

from dotad_release.dictionary_export import export_dictionaries
from dotad_release.models import BuildContext
from dotad_release.workbook_reader import sha256_file


def _dictionary_workbook(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    fields = workbook.active
    fields.title = "Field Dictionary"
    fields.append(
        [
            "source_workbook",
            "source_sheet",
            "field_name",
            "description",
            "source_definition_workbook",
            "source_definition_field",
            "notes",
        ]
    )
    fields.append(
        [
            "release.xlsx",
            "metadata",
            "antibody_id",
            "Exact source field",
            "upstream.xlsx",
            "source_id",
            None,
        ]
    )
    fields.append([None] * 7)
    fields.append(
        [
            "release.xlsx",
            "metadata",
            "antibody_name",
            "Do not map from prose",
            None,
            None,
            "Possible mapping to upstream_name",
        ]
    )

    assays = workbook.create_sheet("Assay Endpoint Dictionary")
    assays.append(
        [
            "Source_Workbook",
            "Original_Field_Key",
            "Harmonized_Endpoint",
            "Risk_Direction",
        ]
    )
    assays.append(["release.xlsx", "hic_rt", "HIC", "higher is riskier"])

    terminology = workbook.create_sheet("Terminology Missingness")
    terminology.append(["term", "definition", "usage_note"])
    terminology.append([None, None, None])
    terminology.append(
        ["antibody-centered", "Organized around antibody records.", "Preferred."]
    )
    workbook.save(path)
    return path


def _context(tmp_path: Path, source: Path) -> BuildContext:
    return BuildContext(
        output_dir=tmp_path / "release",
        core_inputs={"data_dictionary": source},
        sheet_roles={
            "data_dictionary": {
                "field_dictionary": "Field Dictionary",
                "assay_endpoint_dictionary": "Assay Endpoint Dictionary",
                "controlled_vocabulary_source": "Terminology Missingness",
            }
        },
        input_shas={"data_dictionary": sha256_file(source)},
        release_version="2.0.0-draft",
        script_root=tmp_path,
    )


def test_dictionary_exports_preserve_each_nonempty_source_row_and_excel_row(
    tmp_path: Path,
) -> None:
    source = _dictionary_workbook(tmp_path / "dictionary.xlsx")
    results = {
        result.table_name: result
        for result in export_dictionaries(_context(tmp_path, source))
    }

    fields = results["field_dictionary"]
    assert [row["record_id"] for row in fields.rows] == [
        "DICT:field_dictionary:2",
        "DICT:field_dictionary:4",
    ]
    assert [row["source_excel_row"] for row in fields.rows] == [2, 4]
    assert fields.rows[1]["description"] == "Do not map from prose"

    assays = results["assay_endpoint_dictionary"]
    assert [row["record_id"] for row in assays.rows] == [
        "DICT:assay_endpoint_dictionary:2"
    ]
    assert assays.rows[0]["Risk_Direction"] == "higher is riskier"

    controlled = results["controlled_vocabularies"]
    assert [row["record_id"] for row in controlled.rows] == [
        "DICT:terminology_missingness:3"
    ]
    assert controlled.rows[0]["term"] == "antibody-centered"
    payload = json.loads(controlled.rows[0]["source_payload_json"])
    assert payload[1]["value"] == "Organized around antibody records."
    assert controlled.row_definition == (
        "one non-empty source row from Terminology Missingness; "
        "no dedicated controlled-vocabulary sheet is present"
    )


def test_source_field_mapping_uses_only_explicit_field_dictionary_columns(
    tmp_path: Path,
) -> None:
    source = _dictionary_workbook(tmp_path / "dictionary.xlsx")
    results = {
        result.table_name: result
        for result in export_dictionaries(_context(tmp_path, source))
    }

    mapping = results["source_field_mapping"]
    assert mapping.rows == (
        {
            "record_id": "MAP:field_dictionary:2",
            "source_workbook": "release.xlsx",
            "source_sheet": "metadata",
            "source_field": "antibody_id",
            "source_definition_workbook": "upstream.xlsx",
            "source_definition_field": "source_id",
            "field_dictionary_record_id": "DICT:field_dictionary:2",
            "source_excel_row": 2,
            "source_input_sha256": sha256_file(source),
        },
    )


def test_missing_controlled_vocabulary_sheet_creates_semantics_review(
    tmp_path: Path,
) -> None:
    source = _dictionary_workbook(tmp_path / "dictionary.xlsx")
    results = {
        result.table_name: result
        for result in export_dictionaries(_context(tmp_path, source))
    }

    review = results["field_semantics_review_queue"]
    assert len(review.rows) == 1
    assert review.rows[0]["issue_type"] == "MISSING_CONTROLLED_VOCABULARY_SHEET"
    assert review.rows[0]["review_status"] == "NEEDS_REVIEW"
    assert review.rows[0]["fallback_source_sheet"] == "Terminology Missingness"
    assert "dedicated controlled-vocabulary table is absent" in review.rows[0][
        "review_reason"
    ]
