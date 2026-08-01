from __future__ import annotations

import json
from pathlib import Path

import openpyxl

from dotad_release.literature_export import export_literature
from dotad_release.models import BuildContext
from dotad_release.workbook_reader import sha256_file


def _literature_workbook(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    study = workbook.active
    study.title = "Example et al."
    study.append(["antibody_id", "metric", None, "metric"])
    study.append(["DOTAD-001", 1.5, "blank-header", 2.5])
    link = workbook.create_sheet("link")
    link.append(["study", "url"])
    link.append(["Example et al.", "https://example.test"])
    workbook.save(path)
    return path


def _dictionary_workbook(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    sources = workbook.active
    sources.title = "Literature Sources"
    sources.append(
        [
            "source_sheet",
            "pmid",
            "doi",
            "source_url",
            "evidence_type",
        ]
    )
    sources.append(
        [
            "Example et al.",
            "12345678",
            "10.1000/example",
            "https://example.test/source",
            "experimental",
        ]
    )
    workbook.save(path)
    return path


def test_literature_payload_preserves_duplicate_and_blank_headers(
    tmp_path: Path,
) -> None:
    source = _literature_workbook(tmp_path / "literature.xlsx")
    context = BuildContext(
        output_dir=tmp_path / "release",
        core_inputs={"literature": source},
        sheet_roles={"literature": {"excluded": ["link"]}},
        input_shas={"literature": sha256_file(source)},
        release_version="2.0.0-draft",
        script_root=tmp_path,
    )
    result = export_literature(context)
    assert len(result.rows) == 1
    assert result.rows[0]["record_id"] == "LIT:example_et_al:2"
    assert result.rows[0]["literature_record_id"] == "LIT:example_et_al:2"
    assert result.rows[0]["input_file"] == source.name
    assert result.rows[0]["input_file_sha256"] == sha256_file(source)
    assert result.rows[0]["release_version"] == "2.0.0-draft"
    payload = json.loads(result.rows[0]["source_payload_json"])
    assert [item["column_index"] for item in payload] == [1, 2, 3, 4]
    assert [item["header"] for item in payload] == [
        "antibody_id",
        "metric",
        "",
        "metric",
    ]
    assert payload[2]["value"] == "blank-header"
    assert all(row["source_sheet"] != "link" for row in result.rows)


def test_literature_rows_inherit_exact_source_level_citation_metadata(
    tmp_path: Path,
) -> None:
    source = _literature_workbook(tmp_path / "literature.xlsx")
    dictionary = _dictionary_workbook(tmp_path / "dictionary.xlsx")
    context = BuildContext(
        output_dir=tmp_path / "release",
        core_inputs={
            "literature": source,
            "data_dictionary": dictionary,
        },
        sheet_roles={
            "literature": {"excluded": ["link"]},
            "data_dictionary": {"literature_sources": "Literature Sources"},
        },
        input_shas={
            "literature": sha256_file(source),
            "data_dictionary": sha256_file(dictionary),
        },
        release_version="2.0.0-draft",
        script_root=tmp_path,
    )

    result = export_literature(context)
    row = result.rows[0]

    assert row["pmid"] == "12345678"
    assert row["doi"] == "10.1000/example"
    assert row["source_url"] == "https://example.test/source"
    assert row["evidence_type"] == "experimental"
    lineage = json.loads(row["lineage_sources_json"])
    assert {
        (item["input_file"], item["input_sheet"], item["input_excel_row"])
        for item in lineage
    } == {
        (source.name, "Example et al.", 2),
        (dictionary.name, "Literature Sources", 2),
    }
