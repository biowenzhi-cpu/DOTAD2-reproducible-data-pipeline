from __future__ import annotations

import csv
from pathlib import Path

import openpyxl

from dotad_release.metadata_export import export_metadata
from dotad_release.models import BuildContext
from dotad_release.sequence_export import export_sequences
from dotad_release.workbook_reader import sha256_file


SEQUENCE_HEADERS = (
    "antibody_id",
    "antibody_name",
    "VH",
    "VL",
    "hc_protein_sequence",
    "lc_protein_sequence",
    "hc_dna_sequence",
    "lc_dna_sequence",
    "HCDR1",
    "HCDR2",
    "HCDR3",
    "LCDR1",
    "LCDR2",
    "LCDR3",
)

ALLOWED_SEQUENCE_CLASSES = {
    "paired_vh_vl",
    "heavy_only",
    "light_only",
    "full_length_only",
    "cdr_only",
    "mixed_or_other",
    "no_sequence",
}


def _sequence_workbook(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    metadata = workbook.active
    metadata.title = "metadata"
    metadata.append(["antibody_id", "antibody_name"])
    metadata.append(["META-1", "Metadata Name"])
    sequences = workbook.create_sheet("sequences")
    sequences.append(SEQUENCE_HEADERS)
    sequences.append(
        ["PAIR-1", "Pair One", " aa \n", "cc\t"] + [None] * 10
    )
    sequences.append([None] * len(SEQUENCE_HEADERS))
    sequences.append(
        ["HEAVY-1", "Heavy", "AA@*", None] + [None] * 10
    )
    sequences.append(
        ["LIGHT-1", "Light", None, "TT"] + [None] * 10
    )
    sequences.append(
        ["FULL-1", "Full", None, None, "MMMM"] + [None] * 9
    )
    sequences.append(
        ["CDR-1", "CDR", None, None, None, None, None, None, None, None, "CAR"]
        + [None] * 3
    )
    sequences.append(
        [
            "MIXED-1",
            "Mixed",
            "GG",
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            "QQ",
        ]
    )
    sequences.append(
        ["NONE-1", "No sequence"] + [None] * 12
    )
    sequences.append(
        ["PAIR-2", "Pair Duplicate", "AA", "CC"] + [None] * 10
    )
    workbook.save(path)
    return path


def _context(tmp_path: Path, source: Path) -> BuildContext:
    return BuildContext(
        output_dir=tmp_path / "release",
        core_inputs={"metadata_sequences": source},
        sheet_roles={
            "metadata_sequences": {
                "metadata": "metadata",
                "sequences": "sequences",
            }
        },
        input_shas={"metadata_sequences": sha256_file(source)},
        release_version="2.0.0-draft",
        script_root=tmp_path,
    )


def _read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def test_sequence_export_is_source_row_preserving_and_classifies_only_allowed_values(
    tmp_path: Path,
) -> None:
    source = _sequence_workbook(tmp_path / "metadata_sequences.xlsx")

    result = export_sequences(_context(tmp_path, source))

    assert result.fieldnames[: len(SEQUENCE_HEADERS)] == SEQUENCE_HEADERS
    assert len(result.rows) == 8
    assert [row["source_excel_row"] for row in result.rows] == [
        2,
        4,
        5,
        6,
        7,
        8,
        9,
        10,
    ]
    classes = {
        row["antibody_id"]: row["sequence_class"] for row in result.rows
    }
    assert classes == {
        "PAIR-1": "paired_vh_vl",
        "HEAVY-1": "heavy_only",
        "LIGHT-1": "light_only",
        "FULL-1": "full_length_only",
        "CDR-1": "cdr_only",
        "MIXED-1": "mixed_or_other",
        "NONE-1": "no_sequence",
        "PAIR-2": "paired_vh_vl",
    }
    assert set(classes.values()) == ALLOWED_SEQUENCE_CLASSES
    assert result.rows[0]["VH"] == " aa \n"
    assert result.rows[0]["VL"] == "cc\t"
    assert result.rows[0]["record_id"] == "SEQ:sequences:2"
    assert result.rows[1]["record_id"] == "SEQ:sequences:4"
    assert result.rows[0]["sequence_record_id"] == "SEQ:sequences:2"
    assert result.rows[0]["sequence_availability_class"] == "paired_vh_vl"
    assert result.rows[0]["input_file"] == source.name
    assert result.rows[0]["input_file_sha256"] == sha256_file(source)
    assert result.rows[0]["transformation_rule"] == "source_row_preserved"


def test_pair_fingerprint_requires_both_normalized_vh_and_vl(
    tmp_path: Path,
) -> None:
    source = _sequence_workbook(tmp_path / "metadata_sequences.xlsx")

    export_sequences(_context(tmp_path, source))

    fingerprints = _read_tsv(
        tmp_path / "release" / "sequences" / "sequence_fingerprints.tsv"
    )
    by_id = {row["antibody_id"]: row for row in fingerprints}
    assert "CDR-1" not in by_id
    assert by_id["HEAVY-1"]["vh_sha256"] == (
        "d0e56b071a07d070adfa06e2595e1a11eea299a23a18bc2cb67a1cdc7a830c8c"
    )
    assert by_id["HEAVY-1"]["vl_sha256"] == ""
    assert by_id["HEAVY-1"]["pair_sha256"] == ""
    assert by_id["LIGHT-1"]["vh_sha256"] == ""
    assert by_id["LIGHT-1"]["vl_sha256"] != ""
    paired = [row for row in fingerprints if row["pair_sha256"]]
    assert [row["antibody_id"] for row in paired] == ["PAIR-1", "PAIR-2"]
    assert paired[0]["vh_sha256"] == (
        "58bb119c35513a451d24dc20ef0e9031ec85b35bfc919d263e7e5d9868909cb5"
    )
    assert paired[0]["vl_sha256"] == (
        "a56362a10c816abf206d72cb914e2d5ca454eb9c7e744f88b1a1422c379e9942"
    )
    assert paired[0]["pair_sha256"] == (
        "4b2016000a419e860c94b91f5068a1da704e58b980dc6e869ce421d393332a2d"
    )
    assert paired[0]["fingerprint_basis"] == "normalized_vh_vl_sha256"
    assert paired[0]["pair_sha256"] == paired[1]["pair_sha256"]
    assert (
        paired[0]["normalized_vh_vl_pair_sha256"]
        == paired[0]["pair_sha256"]
    )
    assert paired[0]["normalized_vh_sha256"] == paired[0]["vh_sha256"]
    assert paired[0]["normalized_vl_sha256"] == paired[0]["vl_sha256"]
    assert {
        "normalized_heavy_only_sha256",
        "normalized_light_only_sha256",
        "normalized_full_sequence_sha256",
    }.issubset(paired[0])


def test_sequence_qc_reports_properties_without_copying_sequences_into_notes(
    tmp_path: Path,
) -> None:
    source = _sequence_workbook(tmp_path / "metadata_sequences.xlsx")

    export_sequences(_context(tmp_path, source))

    qc_rows = _read_tsv(
        tmp_path / "release" / "sequences" / "sequence_qc.tsv"
    )
    by_id = {row["antibody_id"]: row for row in qc_rows}
    assert len(qc_rows) == 8
    assert by_id["PAIR-1"]["vh_available"] == "true"
    assert by_id["PAIR-1"]["vl_available"] == "true"
    assert by_id["PAIR-1"]["vh_length"] == "2"
    assert by_id["PAIR-1"]["vl_length"] == "2"
    assert by_id["PAIR-1"]["has_whitespace"] == "true"
    assert by_id["HEAVY-1"]["has_illegal_characters"] == "true"
    assert by_id["HEAVY-1"]["illegal_characters"] == "@"
    assert by_id["HEAVY-1"]["has_stop_symbol"] == "true"
    assert "AA@*" not in by_id["HEAVY-1"]["qc_notes"]
    assert by_id["CDR-1"]["duplicate_pair_group_id"] == ""
    assert by_id["PAIR-1"]["duplicate_pair_group_id"].startswith("DUPPAIR:")
    assert (
        by_id["PAIR-1"]["duplicate_pair_group_id"]
        == by_id["PAIR-2"]["duplicate_pair_group_id"]
    )
    assert by_id["PAIR-1"]["duplicate_pair_record_count"] == "2"


def test_alias_sidecar_keeps_metadata_and_sequence_name_occurrences(
    tmp_path: Path,
) -> None:
    source = _sequence_workbook(tmp_path / "metadata_sequences.xlsx")
    context = _context(tmp_path, source)

    export_metadata(context)
    export_sequences(context)
    export_sequences(context)

    aliases = _read_tsv(
        tmp_path / "release" / "metadata" / "antibody_aliases.tsv"
    )
    assert len(aliases) == 9
    assert {row["source_record_id"] for row in aliases} == {
        "META:metadata:2",
        "SEQ:sequences:2",
        "SEQ:sequences:4",
        "SEQ:sequences:5",
        "SEQ:sequences:6",
        "SEQ:sequences:7",
        "SEQ:sequences:8",
        "SEQ:sequences:9",
        "SEQ:sequences:10",
    }
    assert all(row["canonical_name"] == "" for row in aliases)


def test_full_sequence_fingerprint_distinguishes_protein_from_dna(
    tmp_path: Path,
) -> None:
    source = tmp_path / "metadata_sequences.xlsx"
    workbook = openpyxl.Workbook()
    metadata = workbook.active
    metadata.title = "metadata"
    metadata.append(["antibody_id", "antibody_name"])
    sequences = workbook.create_sheet("sequences")
    sequences.append(SEQUENCE_HEADERS)
    sequences.append(
        ["PROTEIN", "Protein", None, None, "ACGT"] + [None] * 9
    )
    sequences.append(
        ["DNA", "DNA", None, None, None, None, "ACGT"] + [None] * 7
    )
    workbook.save(source)

    export_sequences(_context(tmp_path, source))
    rows = {
        row["antibody_id"]: row
        for row in _read_tsv(
            tmp_path / "release" / "sequences" / "sequence_fingerprints.tsv"
        )
    }

    assert (
        rows["PROTEIN"]["normalized_full_sequence_sha256"]
        != rows["DNA"]["normalized_full_sequence_sha256"]
    )
    assert rows["PROTEIN"]["fingerprint_basis"] == (
        "normalized_full_protein_sequence_sha256"
    )
    assert rows["DNA"]["fingerprint_basis"] == (
        "normalized_full_dna_sequence_sha256"
    )
