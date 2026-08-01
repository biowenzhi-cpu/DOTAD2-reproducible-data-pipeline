from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from dotad_release.experimental_export import (
    SOURCE_STUDY_ID,
    export_experimental,
)
from dotad_release.models import BuildContext
from dotad_release.workbook_reader import sha256_file


SHEETS = {
    "GDPa1": {
        "raw": "Assay Data - tidy format",
        "summary": "Assay Data - average ",
    },
    "GDPa2": {
        "raw": "Assay Data - tidy format(2)",
        "summary": "Assay Data - average （2）",
    },
    "GDPa3": {
        "raw": "Assay Data - tidy format（3）",
        "summary": "Assay Data - average（3）",
    },
}

SOURCE_IDS = {
    "GDPa1": "SRC:GINKGO_GDPA1",
    "GDPa2": "SRC:GINKGO_GDPA2_1",
    "GDPa3": "SRC:GINKGO_GDPA3",
}


def _experimental_workbook(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    raw1 = workbook.create_sheet(SHEETS["GDPa1"]["raw"])
    raw1.append(
        [
            "antibody_id",
            "production_batch",
            "technical_replicate",
            "hic_rt",
        ]
    )
    raw1.append(["DOTAD-001", 1, 1, 7.0])
    raw1.append(["DOTAD-001", 1, 2, 9.0])
    summary1 = workbook.create_sheet(SHEETS["GDPa1"]["summary"])
    summary1.append(
        [
            "antibody_id",
            "hic_rt_avg",
            "hic_rt_stddev",
            "hic_rt_replicates",
        ]
    )
    summary1.append(["DOTAD-001", 8.0, 1.41421356237, 2])

    raw2 = workbook.create_sheet(SHEETS["GDPa2"]["raw"])
    raw2.append(
        [
            "antibody_id",
            "production_host",
            "tag",
            "technical_replicate",
            "hic_rt",
        ]
    )
    raw2.append(["DOTAD-002", "CHO", "His", 1, 5.0])
    raw2.append(["DOTAD-002", "HEK", "Fc", 1, 9.0])
    summary2 = workbook.create_sheet(SHEETS["GDPa2"]["summary"])
    summary2.append(
        [
            "antibody_id",
            "production_host",
            "tag",
            "hic_rt_Avg",
            "hic_rt_Count",
        ]
    )
    summary2.append(["DOTAD-002", "CHO", "His", 5.0, 1])
    summary2.append(["DOTAD-002", "HEK", "Fc", 9.0, 1])

    raw3 = workbook.create_sheet(SHEETS["GDPa3"]["raw"])
    raw3.append(["antibody_id", "technical_replicate", "hic_rt"])
    raw3.append(["DOTAD-003", 1, 3.0])
    summary3 = workbook.create_sheet(SHEETS["GDPa3"]["summary"])
    summary3.append(["antibody_id", "hic_rt_avg", "hic_rt_replicates"])
    summary3.append(["DOTAD-003", 3.0, 1])
    workbook.save(path)
    return path


def _context(tmp_path: Path, source: Path) -> BuildContext:
    blocks = {
        block_id: {
            **roles,
            "source_study_id": SOURCE_IDS[block_id],
        }
        for block_id, roles in SHEETS.items()
    }
    return BuildContext(
        output_dir=tmp_path / "release",
        core_inputs={"experimental": source},
        sheet_roles={
            "experimental": {
                "source_study_id": SOURCE_STUDY_ID,
                "blocks": blocks,
            }
        },
        input_shas={"experimental": sha256_file(source)},
        release_version="v2.0.0",
        script_root=tmp_path,
    )


def test_experimental_exports_six_block_layer_files_and_exact_rows(
    tmp_path: Path,
) -> None:
    source = _experimental_workbook(tmp_path / "experimental.xlsx")
    outputs = export_experimental(_context(tmp_path, source))

    assert len(outputs.measurement_results) == 6
    assert {
        result.path.relative_to(tmp_path / "release").as_posix()
        for result in outputs.measurement_results
    } == {
        "evidence/experimental/gdpa1_raw_measurements.tsv",
        "evidence/experimental/gdpa1_summary_measurements.tsv",
        "evidence/experimental/gdpa2_raw_measurements.tsv",
        "evidence/experimental/gdpa2_summary_measurements.tsv",
        "evidence/experimental/gdpa3_raw_measurements.tsv",
        "evidence/experimental/gdpa3_summary_measurements.tsv",
    }
    all_rows = [
        row for result in outputs.measurement_results for row in result.rows
    ]
    assert {row["source_study_id"] for row in all_rows} == set(
        SOURCE_IDS.values()
    )
    assert {row["source_block_id"] for row in all_rows} == {
        "GDPa1",
        "GDPa2",
        "GDPa3",
    }
    assert {
        row["record_layer"] for row in all_rows
    } == {"raw_measurement", "summary_measurement"}
    assert all(row["experimental_record_id"] == row["record_id"] for row in all_rows)
    assert all(row["input_file_sha256"] == sha256_file(source) for row in all_rows)


def test_gdpa2_host_tag_conditions_remain_distinct(tmp_path: Path) -> None:
    source = _experimental_workbook(tmp_path / "experimental.xlsx")
    outputs = export_experimental(_context(tmp_path, source))
    gdpa2_raw = next(
        result
        for result in outputs.measurement_results
        if result.table_name == "gdpa2_raw_measurements"
    )
    assert len(gdpa2_raw.rows) == 2
    assert {
        (row["DOTAD_ID"], row["production_host"], row["tag"])
        for row in gdpa2_raw.rows
    } == {
        ("DOTAD-002", "CHO", "His"),
        ("DOTAD-002", "HEK", "Fc"),
    }
    assert len({row["condition_key"] for row in gdpa2_raw.rows}) == 2


def test_summary_reconstruction_is_block_and_condition_local(
    tmp_path: Path,
) -> None:
    source = _experimental_workbook(tmp_path / "experimental.xlsx")
    outputs = export_experimental(_context(tmp_path, source))
    rows = outputs.summary_reconstruction.rows
    assert len(rows) == 4
    assert {
        (row["source_block_id"], row["verification_status"]) for row in rows
    } == {
        ("GDPa1", "verified_at_reported_precision"),
        ("GDPa2", "verified_exact"),
        ("GDPa3", "verified_exact"),
    }
    gdpa2 = [row for row in rows if row["source_block_id"] == "GDPa2"]
    assert len(gdpa2) == 2
    assert {row["recomputed_summary_value"] for row in gdpa2} == {5.0, 9.0}
    assert {row["recomputed_replicate_count"] for row in gdpa2} == {1}
    assert all(row["aggregation_function"] == "mean" for row in rows)


def test_record_index_and_source_description_keep_three_source_blocks(
    tmp_path: Path,
) -> None:
    source = _experimental_workbook(tmp_path / "experimental.xlsx")
    outputs = export_experimental(_context(tmp_path, source))
    assert len(outputs.record_index.rows) == 9
    assert len(outputs.source_description.rows) == 6
    assert {
        row["source_study_id"] for row in outputs.source_description.rows
    } == set(SOURCE_IDS.values())
    assert {row["source_block_id"] for row in outputs.source_description.rows} == {
        "GDPa1",
        "GDPa2",
        "GDPa3",
    }
    assert {
        row["record_layer"] for row in outputs.source_description.rows
    } == {"raw_measurement", "summary_measurement"}
    assert all(
        row["experimental_record_id"] == row["record_id"]
        for row in outputs.source_description.rows
    )
    assert all(row["source_sheet"] for row in outputs.source_description.rows)
    assert all(row["source_excel_row"] == 1 for row in outputs.source_description.rows)
    assert all(
        row["source_type"] == "within_study_data_block"
        for row in outputs.source_description.rows
    )


def test_blank_reported_summary_cells_are_not_reconstruction_failures(
    tmp_path: Path,
) -> None:
    source = _experimental_workbook(tmp_path / "experimental.xlsx")
    workbook = openpyxl.load_workbook(source)
    summary = workbook[SHEETS["GDPa3"]["summary"]]
    summary["D1"] = "smac_rt_avg"
    summary["D2"] = None
    workbook.save(source)
    outputs = export_experimental(_context(tmp_path, source))
    assert not any(
        row["source_block_id"] == "GDPa3" and row["endpoint"] == "smac_rt"
        for row in outputs.summary_reconstruction.rows
    )


def test_median_summary_is_reconstructed_with_median_not_mean(
    tmp_path: Path,
) -> None:
    source = _experimental_workbook(tmp_path / "experimental.xlsx")
    workbook = openpyxl.load_workbook(source)
    raw = workbook[SHEETS["GDPa3"]["raw"]]
    raw.append(["DOTAD-003", 2, 3.0])
    raw.append(["DOTAD-003", 3, 100.0])
    summary = workbook[SHEETS["GDPa3"]["summary"]]
    summary["B1"] = "hic_rt_median"
    summary["B2"] = 3.0
    summary["C2"] = 3
    workbook.save(source)

    outputs = export_experimental(_context(tmp_path, source))
    row = next(
        item
        for item in outputs.summary_reconstruction.rows
        if item["source_block_id"] == "GDPa3"
        and item["endpoint"] == "hic_rt"
    )

    assert row["aggregation_function"] == "median"
    assert row["recomputed_summary_value"] == 3.0
    assert row["verification_status"] == "verified_exact"


def test_unspecified_summary_condition_does_not_pool_distinct_host_tag_rows(
    tmp_path: Path,
) -> None:
    source = _experimental_workbook(tmp_path / "experimental.xlsx")
    workbook = openpyxl.load_workbook(source)
    summary = workbook[SHEETS["GDPa2"]["summary"]]
    summary.delete_rows(2, summary.max_row)
    summary.append(["DOTAD-002", None, None, 7.0, 2])
    workbook.save(source)

    outputs = export_experimental(_context(tmp_path, source))
    row = next(
        item
        for item in outputs.summary_reconstruction.rows
        if item["source_block_id"] == "GDPa2"
    )

    assert row["recomputed_summary_value"] is None
    assert row["verification_status"] == "partially_verified"
    assert "ambiguous" in row["notes"].lower()


def test_difference_beyond_fixed_tolerance_is_not_reproducible(
    tmp_path: Path,
) -> None:
    source = _experimental_workbook(tmp_path / "experimental.xlsx")
    workbook = openpyxl.load_workbook(source)
    raw = workbook[SHEETS["GDPa3"]["raw"]]
    raw["C2"] = 2.73
    raw.append(["DOTAD-003", 2, 2.74])
    summary = workbook[SHEETS["GDPa3"]["summary"]]
    summary["B2"] = 2.74
    summary["C2"] = 2
    workbook.save(source)

    outputs = export_experimental(_context(tmp_path, source))
    row = next(
        item
        for item in outputs.summary_reconstruction.rows
        if item["source_block_id"] == "GDPa3"
        and item["endpoint"] == "hic_rt"
    )

    assert row["recomputed_summary_value"] == pytest.approx(2.735)
    assert row["verification_status"] == "not_reproducible"


def test_gdpa1_non_titer_endpoint_can_reconstruct_across_batches(
    tmp_path: Path,
) -> None:
    source = _experimental_workbook(tmp_path / "experimental.xlsx")
    workbook = openpyxl.load_workbook(source)
    raw = workbook[SHEETS["GDPa1"]["raw"]]
    raw.append(["DOTAD-001", 2, 1, 11.0])
    summary = workbook[SHEETS["GDPa1"]["summary"]]
    summary["B2"] = 9.0
    summary["C2"] = 2.0
    summary["D2"] = 3
    workbook.save(source)

    outputs = export_experimental(_context(tmp_path, source))
    row = next(
        item
        for item in outputs.summary_reconstruction.rows
        if item["source_block_id"] == "GDPa1"
        and item["endpoint"] == "hic_rt"
    )

    assert row["recomputed_summary_value"] == 9.0
    assert row["recomputed_replicate_count"] == 3
    assert row["verification_status"] == "verified_at_reported_precision"


def test_gdpa1_titer_remains_batch_specific(tmp_path: Path) -> None:
    source = _experimental_workbook(tmp_path / "experimental.xlsx")
    workbook = openpyxl.load_workbook(source)
    raw = workbook[SHEETS["GDPa1"]["raw"]]
    raw["E1"] = "titer"
    raw["E2"] = 100.0
    raw["E3"] = 120.0
    raw.append(["DOTAD-001", 2, 1, 7.0, 300.0])
    summary = workbook[SHEETS["GDPa1"]["summary"]]
    summary["E1"] = "titer_productionbatch1_avg"
    summary["F1"] = "titer_productionbatch1_replicates"
    summary["G1"] = "titer_productionbatch2_avg"
    summary["H1"] = "titer_productionbatch2_replicates"
    summary["E2"] = 110.0
    summary["F2"] = 2
    summary["G2"] = 300.0
    summary["H2"] = 1
    workbook.save(source)

    outputs = export_experimental(_context(tmp_path, source))
    rows = {
        row["endpoint"]: row
        for row in outputs.summary_reconstruction.rows
        if row["source_block_id"] == "GDPa1"
        and row["endpoint"].startswith("titer_")
    }

    assert rows["titer_productionbatch1"]["recomputed_summary_value"] == 110.0
    assert rows["titer_productionbatch1"]["recomputed_replicate_count"] == 2
    assert rows["titer_productionbatch2"]["recomputed_summary_value"] == 300.0
    assert rows["titer_productionbatch2"]["recomputed_replicate_count"] == 1
    assert {
        row["verification_status"] for row in rows.values()
    } == {"verified_at_reported_precision"}


def test_gdpa1_unsuffixed_titer_is_not_pooled_across_batches(
    tmp_path: Path,
) -> None:
    source = _experimental_workbook(tmp_path / "experimental.xlsx")
    workbook = openpyxl.load_workbook(source)
    raw = workbook[SHEETS["GDPa1"]["raw"]]
    raw["E1"] = "titer"
    raw["E2"] = 100.0
    raw["E3"] = 120.0
    raw.append(["DOTAD-001", 2, 1, 7.0, 300.0])
    summary = workbook[SHEETS["GDPa1"]["summary"]]
    summary["E1"] = "titer_avg"
    summary["F1"] = "titer_replicates"
    summary["E2"] = 110.0
    summary["F2"] = 2
    workbook.save(source)

    outputs = export_experimental(_context(tmp_path, source))
    row = next(
        item
        for item in outputs.summary_reconstruction.rows
        if item["source_block_id"] == "GDPa1"
        and item["endpoint"] == "titer"
    )

    assert row["recomputed_summary_value"] is None
    assert row["verification_status"] == "partially_verified"
    assert "production_batch" in row["notes"]


def test_difference_within_fixed_absolute_tolerance_is_verified_rounding(
    tmp_path: Path,
) -> None:
    source = _experimental_workbook(tmp_path / "experimental.xlsx")
    workbook = openpyxl.load_workbook(source)
    raw = workbook[SHEETS["GDPa3"]["raw"]]
    raw["C2"] = 1.0
    summary = workbook[SHEETS["GDPa3"]["summary"]]
    summary["B2"] = 1.000000008
    summary["C2"] = 1
    workbook.save(source)

    outputs = export_experimental(_context(tmp_path, source))
    row = next(
        item
        for item in outputs.summary_reconstruction.rows
        if item["source_block_id"] == "GDPa3"
        and item["endpoint"] == "hic_rt"
    )

    assert row["verification_status"] == "verified_rounding"


def test_replicate_count_must_match_exactly(tmp_path: Path) -> None:
    source = _experimental_workbook(tmp_path / "experimental.xlsx")
    workbook = openpyxl.load_workbook(source)
    summary = workbook[SHEETS["GDPa3"]["summary"]]
    summary["C2"] = 2
    workbook.save(source)

    outputs = export_experimental(_context(tmp_path, source))
    row = next(
        item
        for item in outputs.summary_reconstruction.rows
        if item["source_block_id"] == "GDPa3"
        and item["endpoint"] == "hic_rt"
    )

    assert row["reported_replicate_count"] == 2
    assert row["recomputed_replicate_count"] == 1
    assert row["verification_status"] == "partially_verified"


def test_gdpa1_dls_ph74_summary_maps_to_raw_endpoint(
    tmp_path: Path,
) -> None:
    source = _experimental_workbook(tmp_path / "experimental.xlsx")
    workbook = openpyxl.load_workbook(source)
    raw = workbook[SHEETS["GDPa1"]["raw"]]
    raw["E1"] = "dlskd_ph7.4"
    raw["E2"] = 0.25
    raw["E3"] = 0.25
    summary = workbook[SHEETS["GDPa1"]["summary"]]
    summary["E1"] = "dlskd_pbsph7.4_avg"
    summary["F1"] = "dlskd_pbsph7.4_replicates"
    summary["E2"] = 0.25
    summary["F2"] = 2
    workbook.save(source)

    outputs = export_experimental(_context(tmp_path, source))
    row = next(
        item
        for item in outputs.summary_reconstruction.rows
        if item["source_block_id"] == "GDPa1"
        and item["endpoint"] == "dlskd_pbsph7_4"
    )

    assert row["recomputed_summary_value"] == 0.25
    assert row["recomputed_replicate_count"] == 2
    assert row["verification_status"] == "verified_exact"
    assert outputs.summary_only_endpoints.rows == ()
    assert outputs.experimental_reconstruction_exceptions.rows == ()


def test_gdpa1_dls_ph60_summary_maps_to_raw_endpoint(tmp_path: Path) -> None:
    source = _experimental_workbook(tmp_path / "experimental.xlsx")
    workbook = openpyxl.load_workbook(source)
    raw = workbook[SHEETS["GDPa1"]["raw"]]
    raw["E1"] = "dlskd_ph6.0"
    raw["E2"] = 0.75
    raw["E3"] = 0.75
    summary = workbook[SHEETS["GDPa1"]["summary"]]
    summary["E1"] = "dlskd_hisargph6.0_avg"
    summary["F1"] = "dlskd_hisargph6.0_replicates"
    summary["E2"] = 0.75
    summary["F2"] = 2
    workbook.save(source)

    outputs = export_experimental(_context(tmp_path, source))
    row = next(
        item
        for item in outputs.summary_reconstruction.rows
        if item["source_block_id"] == "GDPa1"
        and item["endpoint"] == "dlskd_hisargph6"
    )

    assert row["recomputed_summary_value"] == 0.75
    assert row["recomputed_replicate_count"] == 2
    assert row["verification_status"] == "verified_exact"
    assert outputs.summary_only_endpoints.rows == ()


@pytest.mark.parametrize(
    ("raw_value", "reported_value"),
    ((72.405, 72.40), (62.305, 62.30)),
)
def test_gdpa1_two_decimal_boundary_is_verified_at_reported_precision(
    tmp_path: Path,
    raw_value: float,
    reported_value: float,
) -> None:
    source = _experimental_workbook(tmp_path / "experimental.xlsx")
    workbook = openpyxl.load_workbook(source)
    raw = workbook[SHEETS["GDPa1"]["raw"]]
    raw["D2"] = raw_value
    raw["D3"] = raw_value
    summary = workbook[SHEETS["GDPa1"]["summary"]]
    summary["B2"] = reported_value
    summary["C2"] = None
    workbook.save(source)

    outputs = export_experimental(_context(tmp_path, source))
    row = next(
        item
        for item in outputs.summary_reconstruction.rows
        if item["source_block_id"] == "GDPa1"
        and item["endpoint"] == "hic_rt"
    )

    assert row["verification_status"] == "verified_at_reported_precision"
    assert row["mean_verification_status"] == "verified_at_reported_precision"
    assert row["stddev_verification_status"] == "not_reported_not_assessed"


def test_gdpa1_two_decimal_difference_above_half_unit_fails(
    tmp_path: Path,
) -> None:
    source = _experimental_workbook(tmp_path / "experimental.xlsx")
    workbook = openpyxl.load_workbook(source)
    raw = workbook[SHEETS["GDPa1"]["raw"]]
    raw["D2"] = 72.4052
    raw["D3"] = 72.4052
    summary = workbook[SHEETS["GDPa1"]["summary"]]
    summary["B2"] = 72.40
    summary["C2"] = None
    workbook.save(source)

    outputs = export_experimental(_context(tmp_path, source))
    row = next(
        item
        for item in outputs.summary_reconstruction.rows
        if item["source_block_id"] == "GDPa1"
        and item["endpoint"] == "hic_rt"
    )

    assert row["verification_status"] == "not_reproducible"


def test_gdpa1_high_precision_polyreactivity_accepts_1e10_difference(
    tmp_path: Path,
) -> None:
    source = _experimental_workbook(tmp_path / "experimental.xlsx")
    workbook = openpyxl.load_workbook(source)
    raw = workbook[SHEETS["GDPa1"]["raw"]]
    raw["E1"] = "polyreactivity_prscore_cho"
    raw["E2"] = 0.1234567891
    raw["E3"] = 0.1234567891
    summary = workbook[SHEETS["GDPa1"]["summary"]]
    summary["E1"] = "polyreactivity_prscore_cho_avg"
    summary["F1"] = "polyreactivity_prscore_cho_replicates"
    summary["E2"] = 0.1234567890
    summary["F2"] = 2
    workbook.save(source)

    outputs = export_experimental(_context(tmp_path, source))
    row = next(
        item
        for item in outputs.summary_reconstruction.rows
        if item["source_block_id"] == "GDPa1"
        and item["endpoint"] == "polyreactivity_prscore_cho"
    )

    assert row["verification_status"] in {"verified_exact", "verified_rounding"}


def test_unreported_stddev_is_not_assessed_or_treated_as_failure(
    tmp_path: Path,
) -> None:
    source = _experimental_workbook(tmp_path / "experimental.xlsx")
    workbook = openpyxl.load_workbook(source)
    summary = workbook[SHEETS["GDPa1"]["summary"]]
    summary["C2"] = None
    workbook.save(source)

    outputs = export_experimental(_context(tmp_path, source))
    row = next(
        item
        for item in outputs.summary_reconstruction.rows
        if item["source_block_id"] == "GDPa1"
        and item["endpoint"] == "hic_rt"
    )

    assert row["verification_status"] == "verified_at_reported_precision"
    assert row["stddev_verification_status"] == "not_reported_not_assessed"
